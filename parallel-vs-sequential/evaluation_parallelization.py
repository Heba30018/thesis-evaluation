#!/usr/bin/env python3
"""
Evaluation runner - the same script for sequential and parallel runs.
Questions are read from test_questions.csv (Question, Answer).

    python3 evaluation_parallelization.py --mode sequential   # 1 thread  -> Evaluation_sequential_results.csv
    python3 evaluation_parallelization.py --mode parallel     # parallel  -> Evaluation_parallel_results.csv
    python3 evaluation_parallelization.py --mode sequential --start 33
Results columns: Question, Answer, Command Match (Match / Not Match),
Execution Status (success / partial / failed), Execution Time (s).
"""

import argparse
import ast
import csv
import json
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import time
from collections import Counter, deque
from datetime import datetime
from pathlib import Path

try:
    import openpyxl          # only needed for .xlsx input
except ImportError:
    openpyxl = None

SCRIPT_DIR = Path(__file__).resolve().parent

# --------------------------------------------------------------------------
# Command normalisation and comparison
# --------------------------------------------------------------------------

# "sonic#", "sonic(config)#", "sonic(conf-if-Management0)#", "leaf1(config-router-bgp)#"
PROMPT_PREFIX = re.compile(r"^\s*[A-Za-z0-9_.\-]+(\([^)]*\))?#\s*")

# Pure navigation lines carry no configuration and are written inconsistently
# in the sheet (some answers have them, some do not), so both sides drop them.
NAVIGATION = {"sonic-cli", "configure terminal", "configure", "conf t", "exit", "end"}


def split_answer(raw):
    """The sheet stores answers as a Python list literal of strings, where one
    string may contain several commands separated by newlines. Some cells are
    plain text instead. Returns a flat list of non-empty lines."""
    if raw is None:
        return []
    if isinstance(raw, (list, tuple)):
        items = list(raw)
    else:
        text = str(raw).strip()
        try:
            value = ast.literal_eval(text)
            items = list(value) if isinstance(value, (list, tuple)) else [str(value)]
        except (ValueError, SyntaxError):
            items = [text]
    lines = []
    for item in items:
        for line in str(item).replace("\\n", "\n").splitlines():
            if line.strip():
                lines.append(line)
    return lines


def normalize_command(line):
    """Removes the CLI prompt prefix, lower-cases, collapses whitespace and
    unifies interface spellings (Ethernet0 / Eth 0 / eth0 -> eth0)."""
    cmd = PROMPT_PREFIX.sub("", line).strip()
    cmd = re.sub(r"\s+", " ", cmd).lower()
    cmd = re.sub(r"\b(ethernet|eth)\s*(\d)", r"eth\2", cmd)
    return cmd


def normalize_commands(lines):
    out = []
    for line in lines:
        cmd = normalize_command(line)
        if cmd and cmd not in NAVIGATION:
            out.append(cmd)
    return out


def compare_commands(expected, generated):
    """Order-independent comparison of the two normalised command lists."""
    exp, gen = Counter(expected), Counter(generated)
    common = sum((exp & gen).values())
    precision = common / sum(gen.values()) if gen else 0.0
    recall = common / sum(exp.values()) if exp else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    if exp and exp == gen:
        label = "Exact"
    elif common > 0:
        label = "Partial"
    else:
        label = "No match"
    missing = list((exp - gen).elements())
    extra = list((gen - exp).elements())
    return label, precision, recall, f1, missing, extra


# --------------------------------------------------------------------------
# Bridge process
# --------------------------------------------------------------------------

class Bridge:
    """Runs `smartclai --netbox-ui-bridge` and timestamps every stdout line
    the moment it arrives (time.monotonic), so stage durations are measured
    on arrival and not when the script gets round to reading them."""

    def __init__(self, binary, env):
        self.proc = subprocess.Popen(
            [str(binary), "--netbox-ui-bridge"],
            cwd=str(binary.parent),          # same as the Electron app
            env=env,
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            bufsize=1,
        )
        self.lines = queue.Queue()
        self.stderr_tail = deque(maxlen=200)
        threading.Thread(target=self._read_stdout, daemon=True).start()
        threading.Thread(target=self._read_stderr, daemon=True).start()

    def _read_stdout(self):
        for line in self.proc.stdout:
            now = time.monotonic()
            line = line.strip()
            if not line:
                continue
            try:
                self.lines.put((now, json.loads(line)))
            except json.JSONDecodeError:
                pass
        self.lines.put((time.monotonic(), None))  # EOF marker

    def _read_stderr(self):
        for line in self.proc.stderr:
            self.stderr_tail.append(line.rstrip("\n"))

    def send(self, obj):
        self.proc.stdin.write(json.dumps(obj, ensure_ascii=False) + "\n")
        self.proc.stdin.flush()
        return time.monotonic()

    def next(self, deadline):
        """Next (time, obj) or (None, None) on timeout. obj None = EOF."""
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            return None, None
        try:
            return self.lines.get(timeout=remaining)
        except queue.Empty:
            return None, None

    def stop(self):
        try:
            self.proc.stdin.close()
        except Exception:
            pass
        try:
            self.proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            self.proc.kill()


def run_generation(bridge, request_id, prompt, timeout_s):
    """Sends the question and collects everything until the bridge's "done"
    for this request. Returns a dict with timestamps and the config action."""
    out = {
        "t_send": None, "t_metrics": None, "t_netbox": None, "t_config": None,
        "t_done": None, "metrics": None, "config": None, "error": None,
        "disambiguation": False,
    }
    out["t_send"] = bridge.send({"type": "ask", "id": request_id, "prompt": prompt})
    deadline = out["t_send"] + timeout_s
    error_grace_until = None

    while True:
        if error_grace_until is not None:
            deadline = min(deadline, error_grace_until)
        t, obj = bridge.next(deadline)
        if t is None:
            if out["error"] is None:
                out["error"] = f"timed out after {timeout_s}s waiting for the generated configuration"
            return out
        if obj is None:
            out["error"] = "bridge process exited"
            return out
        if obj.get("id") != request_id:
            continue

        kind = obj.get("type")
        if kind == "metrics":
            out["t_metrics"], out["metrics"] = t, obj
        elif kind == "error":
            out["error"] = obj.get("message", "unknown error")
            error_grace_until = t + 3.0     # the bridge sends "done" right after
        elif kind == "done":
            out["t_done"] = t
            return out
        elif kind == "json_request":
            message = obj.get("message") or {}
            if message.get("action") == "netbox_result" and out["t_netbox"] is None:
                out["t_netbox"] = t
            elif message.get("type") == "configuration":
                out["t_config"], out["config"] = t, message
            elif message.get("action") == "disambiguation":
                out["disambiguation"] = True


def run_execution(bridge, ask_id, config, timeout_s):
    """Confirms the configuration action (same request the Electron app sends
    when the user presses Execute) and waits for the execution result."""
    exec_id = f"{ask_id}-exec"
    out = {"t_confirm": None, "t_done": None, "payload": None, "error": None}
    out["t_confirm"] = bridge.send({
        "type": "configuration",
        "id": exec_id,
        "accepted": True,
        "action": "config",
        "action_id": config.get("action_id"),
        "targetId": ask_id,
        "commands": config.get("commands") or [],
    })
    deadline = out["t_confirm"] + timeout_s
    while True:
        t, obj = bridge.next(deadline)
        if t is None:
            if out["error"] is None:
                out["error"] = f"timed out after {timeout_s}s waiting for the execution result"
            return out
        if obj is None:
            out["error"] = "bridge process exited"
            return out
        if obj.get("id") != exec_id:
            continue
        if obj.get("type") == "error":
            out["error"] = obj.get("message", "unknown error")
            out["t_done"] = t
            # keep reading: a "done" may follow
            deadline = min(deadline, t + 3.0)
        elif obj.get("type") == "done":
            out["t_done"] = out["t_done"] or t
            out["payload"] = obj.get("payload")
            return out


# --------------------------------------------------------------------------
# Input / output
# --------------------------------------------------------------------------

COLUMNS = ["Question", "Answer", "Command Match", "Execution Status", "Execution Time (s)"]
DEVICE_NAME = re.compile(r"'([^']+)'")


def build_prompt(original, devices):
    s = original.rstrip()
    m = re.search(r"[.?!]+$", s)
    end = m.group(0) if m else "."
    body = s[: m.start()] if m else s
    names = ", ".join(f"'{d.strip()}'" for d in devices)
    return f"{body} for these devices {names}{end}"


def read_questions(path):
    """Reads Question / Answer from a .csv or .xlsx file. In an .xlsx file the
    optional 'Original Question' + 'Devices' columns build the prompt."""
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            records = list(csv.DictReader(f))
    else:
        if openpyxl is None:
            sys.exit("openpyxl is required for .xlsx input:  pip install openpyxl")
        wb = openpyxl.load_workbook(path)
        ws = next(wb[n] for n in wb.sheetnames if n != "Summary")
        header = [str(c.value).strip() if c.value else "" for c in ws[1]]
        records = [dict(zip(header, (c.value for c in row))) for row in ws.iter_rows(min_row=2)]

    rows = []
    for rec in records:
        question = rec.get("Question")
        if not question or not str(question).strip():
            continue
        original, devices_cell = rec.get("Original Question"), rec.get("Devices")
        devices = [d.strip() for d in str(devices_cell).split(",") if d.strip()] if devices_cell else []
        if original and devices:
            prompt = build_prompt(str(original), devices)
        else:
            prompt = str(question)
            devices = DEVICE_NAME.findall(prompt)
        rows.append({"question": str(question), "prompt": prompt,
                     "answer": rec.get("Answer"), "devices": devices})
    return rows


def evaluate(item, gen, exe, dry_run):
    """Returns [Command Match, Execution Status, Execution Time (s)]."""
    config = gen.get("config") or {}

    # Match only when the generated commands equal the expected ones.
    match = "Not Match"
    if config:
        label = compare_commands(normalize_commands(split_answer(item["answer"])),
                                 normalize_commands([str(c) for c in config.get("commands") or []]))[0]
        if label == "Exact":
            match = "Match"

    # success = every requested device ran every command; partial = the backend
    # reported some devices succeeded and others failed; anything else = failed.
    status, exec_time = ("dry run" if dry_run else "failed"), None
    if not dry_run and config and exe is not None and not exe.get("error"):
        payload = exe.get("payload") or {}
        targets = {(t.get("name", "") if isinstance(t, dict) else str(t)).strip().lower()
                   for t in config.get("targets") or []}
        wanted = {d.strip().lower() for d in item["devices"]}
        backend = str(payload.get("status", "")).lower()
        if backend == "partial":
            status = "partial"      # some devices succeeded, others failed
        elif backend == "success" and (not wanted or targets == wanted):
            status = "success"
        timing = payload.get("timing_ms") or {}
        if isinstance(timing.get("execute"), (int, float)):
            exec_time = round(timing["execute"] / 1000.0, 3)
    return [match, status, exec_time]


def save(path, items, results):
    tmp = path.with_suffix(".tmp")
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for item, res in zip(items, results):
            w.writerow([item["question"], item["answer"]] + (res or ["", "", ""]))
    os.replace(tmp, path)


def main():
    parser = argparse.ArgumentParser(
        description="Run the evaluation questions through the SmartCLAI bridge (sequential or parallel).",
    )
    parser.add_argument("--mode", choices=["sequential", "parallel"], default="parallel",
                        help="sequential = one device at a time (1 thread); parallel = devices at the same time")
    parser.add_argument("--input", type=Path, default=SCRIPT_DIR / "test_questions.csv",
                        help="questions file (.csv or .xlsx) with Question and Answer columns")
    parser.add_argument("--output", type=Path, default=None,
                        help="results file (default: Evaluation_<mode>_results.csv)")
    parser.add_argument("--binary", type=Path, default=SCRIPT_DIR.parent / "build" / "smartclai")
    parser.add_argument("--start", type=int, default=1, help="first question to run (1-based)")
    parser.add_argument("--limit", type=int, default=None, help="run at most N questions")
    parser.add_argument("--dry-run", action="store_true",
                        help="generate and compare commands only; never execute on devices")
    parser.add_argument("--concurrency", type=int, default=None,
                        help="parallel mode only: max devices at the same time (default 10)")
    parser.add_argument("--generation-timeout", type=float, default=180.0)
    parser.add_argument("--execution-timeout", type=float, default=900.0)
    parser.add_argument("--only-failed", action="store_true",
                        help="run only questions whose result in the output file is not 'success'")
    parser.add_argument("--pause", type=float, default=0.0, help="seconds to wait between questions")
    args = parser.parse_args()

    binary = args.binary.resolve()
    if not binary.exists():
        sys.exit(f"Binary not found: {binary}  (build first, or pass --binary)")
    if not args.input.exists():
        sys.exit(f"Input file not found: {args.input}")
    output = args.output or SCRIPT_DIR / f"Evaluation_{args.mode}_results.csv"

    all_items = read_questions(args.input)
    if not all_items:
        sys.exit(f"{args.input}: no questions found (needs Question and Answer columns)")

    # Keep results already in the output file for questions not run this time.
    results = [None] * len(all_items)
    if output.exists():
        if output.resolve() == args.input.resolve():
            backup = output.with_name(output.stem + "_backup.csv")
            if not backup.exists():
                shutil.copy2(output, backup)
        with open(output, newline="", encoding="utf-8-sig") as f:
            old = {r.get("Question"): r for r in csv.DictReader(f)}
        for i, item in enumerate(all_items):
            r = old.get(item["question"])
            if r and r.get("Command Match"):
                results[i] = [r.get("Command Match"), r.get("Execution Status"), r.get("Execution Time (s)")]

    first = args.start - 1
    last = len(all_items) if not args.limit else min(len(all_items), first + args.limit)
    todo = list(range(first, last))
    if args.only_failed:
        todo = [i for i in todo if not results[i] or results[i][1] != "success"]
    error_log = output.with_name(output.stem.replace("_results", "") + "_errors.log")

    # Prewarm is always off: SSH connect, commands and disconnect of every
    # device all happen inside the measured execution time.
    env = dict(os.environ, SMARTCLAI_EMIT_METRICS="1", SMARTCLAI_DISABLE_PREWARM="1")
    if args.mode == "sequential":
        env["SMARTCLAI_SSH_MAX_CONCURRENCY"] = "1"
        concurrency = 1
    elif args.concurrency is not None:
        env["SMARTCLAI_SSH_MAX_CONCURRENCY"] = str(args.concurrency)
        concurrency = args.concurrency
    else:
        env.pop("SMARTCLAI_SSH_MAX_CONCURRENCY", None)
        concurrency = "default (10)"

    print(f"Mode: {args.mode}   concurrency: {concurrency}   questions: {len(todo)}"
          f"   {'DRY RUN' if args.dry_run else 'execute on devices'}")
    print(f"Input:  {args.input}\nOutput: {output}\n")
    if not args.dry_run:
        print("WARNING: commands WILL be executed on the devices. Ctrl+C now to stop.\n")
        time.sleep(3)

    bridge = Bridge(binary, env)
    run_tag = datetime.now().strftime("%H%M%S")
    done = 0
    try:
        for n, i in enumerate(todo, start=1):
            item = all_items[i]
            ask_id = f"{args.mode[:3]}-{run_tag}-{i + 1}"
            label = f"[{n}/{len(todo)}] question {i + 1} ({len(item['devices'])} devices)"
            print(f"{label}: generating ...", flush=True)

            gen = run_generation(bridge, ask_id, item["prompt"], args.generation_timeout)
            exe = None
            if gen.get("config") and not args.dry_run:
                print(f"{label}: executing ...", flush=True)
                exe = run_execution(bridge, ask_id, gen["config"], args.execution_timeout)

            results[i] = evaluate(item, gen, exe, args.dry_run)

            # Why did it fail? Print it and append it to the errors log.
            reasons = []
            if gen.get("error"):
                reasons.append(f"generation: {gen['error']}")
            if exe and exe.get("error"):
                reasons.append(f"execution: {exe['error']}")
            for d in ((exe or {}).get("payload") or {}).get("device_results") or []:
                if not d.get("success"):
                    reasons.append(f"{d.get('name')}: {d.get('error')}")
            if gen.get("config"):
                generated = gen["config"].get("commands") or []
            else:
                generated = []
            if reasons:
                for line in reasons:
                    print(f"    ! {line}")
                with open(error_log, "a", encoding="utf-8") as f:
                    f.write(f"[{datetime.now().isoformat(timespec='seconds')}] question {i + 1} "
                            f"({args.mode}) {results[i][0]} | {results[i][1]}\n"
                            f"  Q: {item['prompt']}\n"
                            f"  generated: {generated}\n"
                            + "".join(f"  ! {line}\n" for line in reasons) + "\n")
            save(output, all_items, results)
            done += 1
            match, status, secs_ = results[i]
            print(f"{label}: {match} | {status} | {secs_}s")

            if bridge.proc.poll() is not None:
                print("\nThe bridge process stopped. Last stderr lines:")
                print("\n".join(list(bridge.stderr_tail)[-30:]))
                break
            if args.pause:
                time.sleep(args.pause)
    except KeyboardInterrupt:
        print("\nInterrupted - results so far are saved.")
    finally:
        bridge.stop()
        save(output, all_items, results)

    print(f"\nDone. {done} question(s) written to {output}")


if __name__ == "__main__":
    main()
